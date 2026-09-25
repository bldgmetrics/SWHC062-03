"""Savings-workbook chain for the 2026-09-15 M3/M5 rerun.
Refresh hourly_results -> Deer Peak - Electric -> -all files -> off-hour merge
-> fill 20260915 template -> write Off-hour tab. Excel recalc done after."""
import csv as csvmod
import os
import shutil
import subprocess
import sys
import time

REPO = r"c:\dev\SWHC062-03"
PY = os.path.join(REPO, r".venv\Scripts\python.exe")
PP = os.path.join(REPO, "postprocess")
MEAS = os.path.join(REPO, r"commercial measures\SWHC062-03 Occupancy Fan Controller")
WB_OUT = "SWHC062_Energy_Savings_Calculations_20260915_updated.xlsx"


def log(m):
    print(f"[{time.strftime('%H:%M:%S')}] {m}", flush=True)


def step(name, args):
    t0 = time.time()
    r = subprocess.run([PY] + args, cwd=PP, stdout=subprocess.PIPE,
                       stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace")
    log(f"[{name}] rc={r.returncode} {time.time()-t0:.0f}s | " +
        " / ".join(r.stdout.strip().splitlines()[-1:]))
    if r.returncode != 0:
        print(r.stdout[-2000:], flush=True)
        sys.exit(1)


# fresh hourly_results for the off-hour merge
dst = os.path.join(PP, "hourly_results")
if os.path.isdir(dst):
    shutil.rmtree(dst)
shutil.copytree(os.path.join(MEAS, "hourly_results"), dst)
log("hourly_results refreshed")

step("make_deer_peak_electric", ["make_deer_peak_electric.py"])
step("build_all_files", ["build_all_files.py"])
step("make_measure_off_hour", ["make_measure_off_hour.py"])
step("com_energy_savings_calc", ["com_energy_savings_calc.py"])

# write Measure Off-hour tab (cols A:J, rows 4+) from the fresh merge
import openpyxl
wb = openpyxl.load_workbook(os.path.join(PP, WB_OUT))
ws = wb["Measure Off-hour"]
with open(os.path.join(PP, "Measure Off-hour2.csv"), newline="") as f:
    data = list(csvmod.reader(f))[1:]
for rr in range(4, ws.max_row + 1):
    for cc in range(1, 11):
        ws.cell(row=rr, column=cc).value = None
for i, row in enumerate(data, start=4):
    for cc in range(10):
        v = row[cc]
        try:
            v = float(v) if cc >= 4 else v
        except ValueError:
            pass
        ws.cell(row=i, column=cc + 1).value = v
wb.save(os.path.join(PP, WB_OUT))
log(f"Measure Off-hour tab: {len(data)} rows -> {WB_OUT}")
log("WORKBOOK CHAIN COMPLETE")
